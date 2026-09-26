"""Compare resolved endpoints against the sponsor's worked example.

The sponsor spreadsheet is a secondary source (handoff I-1/I-3): it is read only to check our
output and never feeds geometry into the pipeline.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from gridlock.ingestion.documents.plans import raw_dir
from gridlock.models.domain import EndpointMatchStatus, Point, Project
from gridlock.settings.loader import ConfigBundle

from .matching import distance_km
from .names import name_score, normalize_name

ENDPOINT_COLUMNS = (("name_a", "lat_a", "lon_a"), ("name_b", "lat_b", "lon_b"))


@dataclass(frozen=True)
class OracleRow:
    sponsor_id: str
    project_id: str
    sheet_endpoint: str
    sheet_point: Point | None
    our_status: str
    our_feature: str | None
    distance_km: float | None

    def as_list(self) -> list[object]:
        sheet = f"{self.sheet_point.lat:.5f},{self.sheet_point.lon:.5f}" if self.sheet_point else ""
        distance = f"{self.distance_km:.2f}" if self.distance_km is not None else ""
        return [self.sponsor_id, self.project_id, self.sheet_endpoint, sheet, self.our_status, self.our_feature or "", distance]


HEADER = ["sponsor_id", "project_id", "sheet_endpoint", "sheet_lat_lon", "our_status", "our_feature", "distance_km"]


def oracle_check(
    bundle: ConfigBundle, repository_root: Path, output_path: Path | None = None, resolved_path: Path | None = None
) -> list[OracleRow]:
    oracle = bundle.root.oracle
    if oracle is None:
        raise ValueError("no oracle is configured")
    rules = bundle.root.resolution
    resolved_path = resolved_path or repository_root / bundle.root.paths.normalized_dir / "projects_resolved.json"
    projects = {item["id"]: Project.model_validate(item) for item in json.loads(resolved_path.read_text(encoding="utf-8"))}
    sheet = load_workbook(raw_dir(bundle, repository_root) / oracle.file, read_only=True, data_only=True)[oracle.sheet]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value) for value in next(rows)]

    results: list[OracleRow] = []
    for values in rows:
        record = dict(zip(header, values))
        sponsor_id = str(record.get("project_id") or "")
        project_id = oracle.project_ids.get(sponsor_id)
        if project_id is None:
            continue
        project = projects[project_id]
        for name_column, lat_column, lon_column in ENDPOINT_COLUMNS:
            sheet_name = str(record.get(name_column) or "")
            if not sheet_name:
                continue
            lat, lon = record.get(lat_column), record.get(lon_column)
            sheet_point = Point(lat=float(lat), lon=float(lon)) if lat is not None and lon is not None else None
            target = normalize_name(sheet_name, rules)
            matches = project.geometry_resolution.matches if project.geometry_resolution else []
            ours = max(matches, key=lambda match: name_score(target, normalize_name(match.endpoint, rules)), default=None)
            if ours is None or name_score(target, normalize_name(ours.endpoint, rules)) < rules.name_match_threshold:
                results.append(OracleRow(sponsor_id, project_id, sheet_name, sheet_point, "not_in_title", None, None))
                continue
            resolved = ours.status in {EndpointMatchStatus.MATCHED, EndpointMatchStatus.OVERRIDE} and ours.point is not None
            distance = distance_km(sheet_point, ours.point) if resolved and sheet_point else None
            label = ours.feature_name or ("human-verified" if ours.status is EndpointMatchStatus.OVERRIDE else "")
            feature = f"{label} ({ours.feature_id})" if ours.feature_id else label or None
            results.append(OracleRow(sponsor_id, project_id, sheet_name, sheet_point, ours.status.value, feature, distance))

    output_path = output_path or repository_root / bundle.root.paths.review_dir / "oracle_check.csv"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as output:
        writer = csv.writer(output)
        writer.writerow(HEADER)
        writer.writerows(row.as_list() for row in results)
    return results
